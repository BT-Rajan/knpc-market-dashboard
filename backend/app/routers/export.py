import csv
import io
import re

import pandas as pd
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl.styles import Font
from sqlalchemy import inspect
from sqlalchemy.orm import Session

from app.db import get_db
from app.auth import get_current_admin
from app.models import Item, Source, PriceHistory, NewsItem, ScrapeLog, ScrapeSetting
from app.services import item_price_export_rows

router = APIRouter(prefix="/api/admin/export", tags=["export"], dependencies=[Depends(get_current_admin)])

TABLES = {
    "items": Item,
    "sources": Source,
    "price_history": PriceHistory,
    "news_items": NewsItem,
    "scrape_log": ScrapeLog,
    "scrape_settings": ScrapeSetting,
}

_INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")


@router.get("/prices.xlsx")
def export_prices_xlsx(db: Session = Depends(get_db)):
    """One sheet per active crude benchmark and product, named by item code
    (always short and unique, unlike the long display names) with the full
    name and unit as a title inside the sheet. Crude sheets are daily; product
    sheets are week-ending, matching how each is actually re-priced."""
    items = db.query(Item).filter(Item.active == True).order_by(Item.category, Item.name).all()  # noqa: E712

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for item in items:
            is_product = item.category == "Products"
            date_col = "Week ending" if is_product else "Date"
            price_col = f"Price ({item.unit})"
            rows = item_price_export_rows(db, item)
            df = pd.DataFrame(
                [{date_col: r["date"], price_col: round(r["price"], 4)} for r in rows],
                columns=[date_col, price_col],
            )

            sheet_name = _INVALID_SHEET_CHARS.sub("", item.code)[:31] or f"item{item.id}"
            df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=3)

            ws = writer.sheets[sheet_name]
            ws["A1"] = item.name
            ws["A1"].font = Font(bold=True, size=12)
            cadence = "week-ending" if is_product else "daily"
            ws["A2"] = f"{item.category} · {item.unit} · {cadence}"
            ws["A2"].font = Font(italic=True, size=9, color="808080")
            ws.column_dimensions["A"].width = 16
            ws.column_dimensions["B"].width = 16
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=price_export.xlsx"},
    )


@router.get("/tables")
def list_tables():
    """Which raw tables can be exported."""
    return {"tables": list(TABLES.keys())}


@router.get("/tables/{table_name}.csv")
def export_table_csv(table_name: str, db: Session = Depends(get_db)):
    model = TABLES.get(table_name)
    if not model:
        raise HTTPException(status_code=404, detail=f"Unknown table '{table_name}'")

    columns = [c.key for c in inspect(model).columns]
    rows = db.query(model).all()

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(columns)
    for row in rows:
        writer.writerow([getattr(row, col) for col in columns])
    buffer.seek(0)

    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={table_name}.csv"},
    )
